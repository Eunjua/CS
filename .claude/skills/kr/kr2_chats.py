#!/usr/bin/env python3
"""KR2 가중치용 — 채널톡 export에서 채널별·주차별 채팅 상담 수를 센다.

  python3 kr2_chats.py 채널톡 --out <저장소 밖 JSON>
  python3 kr2_chats.py --selftest

- 채널: `Manager data` 시트 channelId — 201411 = 케어아카데미, 87358 = 케어파트너 (/월요일과 같은 기준)
- 주차: managedAt 기준 월~일 (VOC 대시보드 v2와 같은 기준). 키는 그 주 월요일 YYYY-MM-DD
- 채팅만: mediumType=phone은 뺀다 (케어파트너 export에 전화가 섞여 있다)
- 같은 상담 id는 한 번만 센다 (파일이 겹쳐도)
- 출력은 건수뿐이라 개인정보가 없지만, export와 같이 저장소 밖에 둔다
"""
import collections
import datetime as dt
import glob
import json
import os
import sys
import unicodedata

CHANNELS = {"201411": "academy", "87358": "partner"}
NAMES = {"academy": "케어아카데미", "partner": "케어파트너"}


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v or "").strip().replace("T", " ")
    if not s:
        return None
    try:
        return dt.date.fromisoformat(s[:10])
    except ValueError:
        return None


def monday(d):
    return d - dt.timedelta(days=d.weekday())


def file_channel(sheets):
    md = sheets.get("Manager data")
    if md is None:
        return None
    rows = md.iter_rows(values_only=True)
    head = list(next(rows, []))
    if "channelId" not in head:
        return None
    i = head.index("channelId")
    for r in rows:
        ch = str(r[i] or "").strip()
        if ch:
            return CHANNELS.get(ch, "unknown:" + ch)
    return None


def count(paths):
    from openpyxl import load_workbook

    seen, out, notes = set(), collections.defaultdict(lambda: collections.Counter()), []
    for p in paths:
        wb = load_workbook(p, read_only=True, data_only=True)
        sheets = {unicodedata.normalize("NFC", ws.title): ws for ws in wb.worksheets}
        ch = file_channel(sheets)
        name = os.path.basename(p)
        if ch not in NAMES:
            notes.append(f"{name}: 채널을 못 가렸어요({ch}) — 건너뜀")
            continue
        uc = sheets.get("UserChat")
        if uc is None:
            notes.append(f"{name}: UserChat 시트 없음 — 건너뜀")
            continue
        rows = uc.iter_rows(values_only=True)
        head = list(next(rows, []))
        idx = {h: i for i, h in enumerate(head)}
        if "id" not in idx or "managedAt" not in idx:
            notes.append(f"{name}: id·managedAt 칸 없음 — 건너뜀")
            continue
        phone = dup = nodate = n = 0
        for r in rows:
            cid = r[idx["id"]]
            if not cid:
                continue
            if cid in seen:
                dup += 1
                continue
            seen.add(cid)
            if "mediumType" in idx and str(r[idx["mediumType"]] or "").lower() == "phone":
                phone += 1
                continue
            d = to_date(r[idx["managedAt"]])
            if not d:
                nodate += 1
                continue
            out[monday(d).isoformat()][ch] += 1
            n += 1
        notes.append(f"{name}: {NAMES[ch]} 채팅 {n}건 (전화 제외 {phone} · 중복 {dup} · 날짜 없음 {nodate})")
    return {w: dict(c) for w, c in sorted(out.items())}, notes


def selftest():
    assert monday(dt.date(2026, 10, 12)) == dt.date(2026, 10, 12)
    assert monday(dt.date(2026, 10, 18)) == dt.date(2026, 10, 12)
    assert to_date("2026-10-07 09:12:00") == dt.date(2026, 10, 7)
    assert to_date(dt.datetime(2026, 10, 7, 9)) == dt.date(2026, 10, 7)
    assert to_date("") is None
    print("kr2_chats.py selftest ok")


def main(argv):
    if argv[:1] == ["--selftest"]:
        return selftest()
    if "--out" not in argv:
        sys.exit(__doc__)
    i = argv.index("--out")
    out = argv[i + 1]
    folders = argv[:i] + argv[i + 2:]
    paths = sorted(p for f in folders for p in glob.glob(os.path.join(f, "*.xlsx")) if not os.path.basename(p).startswith("~$"))
    if not paths:
        sys.exit("export xlsx가 없어요 — 채널톡/ 폴더에 넣어 주세요")
    weeks, notes = count(paths)
    print("\n".join(notes))
    for w, c in weeks.items():
        print(f"{w} 주  케어아카데미 {c.get('academy', 0):>4}  케어파트너 {c.get('partner', 0):>4}")
    json.dump(weeks, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"→ {out}")


if __name__ == "__main__":
    main(sys.argv[1:])
