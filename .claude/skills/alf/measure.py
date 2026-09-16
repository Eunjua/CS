"""ALF 회차 지표 — 채널톡 export에서 분모·봇완결률·이탈률을 직접 센다.

사용법:
  python3 measure.py cs-triage/상담내역/*.xlsx --out <저장소 밖 폴더>
  python3 measure.py --selftest

- 지표는 이 스크립트 출력에서만 가져온다. 손으로 세지 않는다.
- 판정 정의는 reference/metrics.md. 2026-09-11 개정(15차~) 기준이다.
- 건별 CSV에는 상담 링크·태그가 들어간다 → 저장소 밖에만 쓴다(--out 필수).
"""
import collections
import csv
import glob
import os
import re
import sys
import unicodedata

# 이탈 판정 정규식 — 15차(2026-09-11)에 쓴 그대로. 바꾸면 과거 회차와 계열이 끊긴다.
PHONE = r'전화번호를 (먼저 )?알려주시겠어요|전화번호를 다시'
RESET = r'어떤 도움이 필요하신가요|궁금하신 내용을 말씀해 주시면'
ASK = r'\?|과정을 말씀해 주세요'
CLOSE = r'더 도와드릴 것이 있을까요|다른 도움이 필요하신|궁금하신 점|궁금한 점'
URL = r'https?://\S+'

ASSIGN_COLS = ['assigneeId', 'firstAssigneeId', 'managerIds']


def is_bounce(last_bot_text):
    """미배정 + 마지막 발화가 봇일 때, 그 발화가 고객을 세워둔 채 끝났는가."""
    t = re.sub(URL, '', last_bot_text or '')
    if re.search(PHONE, t) or re.search(RESET, t):
        return True
    return bool(re.search(ASK, t)) and not re.search(CLOSE, t)


def load(path):
    """export 한 개에서 상담별 dict를 뽑는다. id → 상담."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = {unicodedata.normalize('NFC', ws.title): ws for ws in wb.worksheets}
    uc = sheets.get('UserChat')
    if uc is None:
        sys.exit(f'{os.path.basename(path)}: UserChat 시트가 없습니다')

    rows = list(uc.iter_rows(values_only=True))
    idx = {h: i for i, h in enumerate(rows[0])}
    chats = {}
    for r in rows[1:]:
        cid = r[idx['id']]
        if not cid:
            continue
        chats[cid] = {
            'id': cid,
            'url': r[idx['url']] or '',
            'createdAt': str(r[idx['createdAt']] or ''),
            'tags': [t.strip() for t in (r[idx['tags']] or '').split(',') if t.strip()],
            'alf': str(r[idx['alfTriggered']]).lower() == 'true',
            'assigned': any(r[idx[c]] not in (None, '') for c in ASSIGN_COLS),
            'bot_msgs': 0,
            'last_type': '',
            'last_bot_text': '',
        }

    md = sheets.get('Message data')
    if md is not None:
        mrows = list(md.iter_rows(values_only=True))
        mi = {h: i for i, h in enumerate(mrows[0])}
        msgs = collections.defaultdict(list)
        for n, r in enumerate(mrows[1:]):
            cid = r[mi['chatId']]
            if cid in chats:
                # 상담원 내부 메모(isPrivate)는 고객에게 안 보이므로 마지막 발화에서 뺀다.
                if str(r[mi['isPrivate']]).lower() == 'true':
                    continue
                msgs[cid].append((str(r[mi['createdAt']] or ''), n, r[mi['personType']], r[mi['plainText']] or ''))
        for cid, ms in msgs.items():
            ms.sort()
            chats[cid]['bot_msgs'] = sum(1 for m in ms if m[2] == 'bot')
            chats[cid]['last_type'] = ms[-1][2]
            if ms[-1][2] == 'bot':
                chats[cid]['last_bot_text'] = ms[-1][3]
    return chats


def measure(chats):
    """상담 dict들 → 지표. 판정 결과를 각 상담에 써 넣는다."""
    total = len(chats)
    for c in chats.values():
        c['done'] = c['alf'] and not c['assigned']
        c['bounce'] = c['done'] and c['last_type'] == 'bot' and is_bounce(c['last_bot_text'])
    alf = [c for c in chats.values() if c['alf']]
    done = [c for c in alf if c['done']]
    bounce = [c for c in alf if c['bounce']]
    tags = collections.Counter(t for c in alf for t in c['tags'] if t.startswith('AI상담/'))
    return {
        'total': total,
        'alf': len(alf),
        'no_alf': total - len(alf),
        'done': len(done),
        'done_rate': len(done) / len(alf) if alf else 0,
        'bounce': len(bounce),
        'bounce_rate': len(bounce) / len(alf) if alf else 0,
        'bot_msgs': sum(c['bot_msgs'] for c in chats.values()),
        'ai_tags': tags,
    }


def report(m, chats):
    print(f"\n분모 산출: 전체 {m['total']} → alf발동 {m['alf']}(미발동 {m['no_alf']})"
          f" → 봇완결(미배정) {m['done']}")
    print(f"결과: 완결률 {m['done_rate']:.1%}({m['done']}/{m['alf']})"
          f" · 이탈률 {m['bounce_rate']:.1%}({m['bounce']}건, 행동 기준)")

    print(f"\n검증: alf발동 {m['alf']} + 미발동 {m['no_alf']} = {m['alf'] + m['no_alf']}"
          f" (총 상담 {m['total']}) → {'OK' if m['alf'] + m['no_alf'] == m['total'] else '❌ 불일치'}")

    if m['bot_msgs'] == 0:
        print("\n⚠️ 이 export에는 봇 발화(personType='bot')가 없습니다."
              " 말투·정확도 직접 채점은 못 하고 이관 적정성 축으로만 판정하세요."
              " 이탈률도 셀 수 없으니 결과에 그 한계를 밝히세요.")
    else:
        print(f"\n봇 발화 {m['bot_msgs']}건 — 말투·정확도 직접 채점 가능")

    if m['ai_tags']:
        print('\nAI상담/* 태그:', ' · '.join(f'{k.split("/")[-1]} {v}' for k, v in m['ai_tags'].most_common()))
    else:
        print('\nAI상담/* 태그: 0건 — 태그로 놓침을 셀 수 없습니다. 대화를 읽고 판정하세요.')

    days = collections.Counter(c['createdAt'][:10] for c in chats.values() if c['createdAt'])
    if days:
        last = max(days)
        print(f"\n⚠️ 마지막 날짜({last}) {days[last]}건은 export 시점에 진행 중일 수 있습니다 — 판정 시 표시해 두세요.")


def write_csv(chats, out):
    """건별 판정표. 상담 링크가 들어가므로 저장소 밖에만 쓴다."""
    repo = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
    if os.path.abspath(out).startswith(repo + os.sep):
        sys.exit(f'--out은 저장소 밖이어야 합니다(상담 링크 포함). 지금 값: {out}')
    os.makedirs(out, exist_ok=True)
    path = os.path.join(out, 'alf_건별.csv')
    cols = ['id', 'url', 'createdAt', 'tags', 'alfTriggered', '배정', '봇완결', '이탈',
            '봇발화수', '마지막발화', '마지막봇발화']
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(cols)
        for c in sorted(chats.values(), key=lambda x: x['createdAt']):
            w.writerow([c['id'], c['url'], c['createdAt'], ', '.join(c['tags']),
                        'Y' if c['alf'] else 'N', 'Y' if c['assigned'] else '',
                        'Y' if c['done'] else '', 'Y' if c['bounce'] else '',
                        c['bot_msgs'], c['last_type'], (c['last_bot_text'] or '')[:60].replace('\n', ' ')])
    print(f'\n건별 판정표 {len(chats)}행 → {path}')
    return path


def selftest():
    assert is_bounce('전화번호를 먼저 알려주시겠어요?')
    assert is_bounce('어떤 도움이 필요하신가요')
    assert is_bounce('어떤 과정을 말씀해 주세요')
    assert not is_bounce('신청은 https://academy.carepartner.kr/a?b=1 에서 하시면 됩니다')  # URL 속 물음표
    assert not is_bounce('환불은 7일 이내 가능합니다. 더 도와드릴 것이 있을까요?')         # 마무리 질문
    assert not is_bounce('발급비는 93,000원입니다.')

    chats = {
        'a': {'alf': True, 'assigned': False, 'last_type': 'bot', 'last_bot_text': '전화번호를 다시 알려주세요', 'tags': [], 'bot_msgs': 2, 'createdAt': '2026-09-01'},
        'b': {'alf': True, 'assigned': False, 'last_type': 'user', 'last_bot_text': '', 'tags': ['AI상담/오안내'], 'bot_msgs': 3, 'createdAt': '2026-09-01'},
        'c': {'alf': True, 'assigned': True, 'last_type': 'manager', 'last_bot_text': '', 'tags': [], 'bot_msgs': 1, 'createdAt': '2026-09-01'},
        'd': {'alf': False, 'assigned': True, 'last_type': 'manager', 'last_bot_text': '', 'tags': [], 'bot_msgs': 0, 'createdAt': '2026-09-01'},
    }
    m = measure(chats)
    assert (m['total'], m['alf'], m['no_alf'], m['done'], m['bounce']) == (4, 3, 1, 2, 1), m
    assert abs(m['done_rate'] - 2 / 3) < 1e-9 and abs(m['bounce_rate'] - 1 / 3) < 1e-9
    assert m['ai_tags']['AI상담/오안내'] == 1
    print('measure.py selftest ok')


def main(argv):
    if argv[:1] == ['--selftest']:
        return selftest()
    out = None
    if '--out' in argv:
        i = argv.index('--out')
        out = argv[i + 1] if len(argv) > i + 1 else sys.exit('--out 뒤에 폴더를 적어주세요')
        argv = argv[:i] + argv[i + 2:]
    paths = [p for a in argv for p in sorted(glob.glob(a))]
    if not paths:
        sys.exit(__doc__)

    chats = {}
    for p in paths:
        got = load(p)
        dup = len(set(got) & set(chats))
        chats.update(got)  # id로 중복 제거(union)
        print(f'{os.path.basename(p)[:40]}: {len(got)}건' + (f' (중복 {dup} 제외)' if dup else ''))

    m = measure(chats)
    report(m, chats)
    if out:
        write_csv(chats, out)
    else:
        print('\n(건별 판정표가 필요하면 --out <저장소 밖 폴더>를 붙이세요)')


if __name__ == '__main__':
    main(sys.argv[1:])
